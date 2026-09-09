"""Real OOXML preservation and rendered editorial regression in isolated Docker.

The model/reviewer responses are fixed fixtures. The DOCX parser, revision
controller, package comparisons, LibreOffice and PDF rendering really execute.
Optional --source imports a supplied DOCX in addition to the controlled fixture.
"""
from __future__ import annotations
import argparse
from pathlib import Path
import shutil
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from cumcm_harness.common import Blocked,file_hash,write_json,read_json
from cumcm_harness.sandbox import Executor,Limits


def worker(input_dir,out):
    import io,zipfile,subprocess
    from PIL import Image
    import pymupdf as fitz
    from lxml import etree as ET
    from docx import Document
    from docx.shared import Cm,Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from cumcm_harness.controller import DEFAULT_CONFIG
    from cumcm_harness.entry_inputs import initialize
    from cumcm_harness.entry_documents import read_document,W,M,sha_text
    from cumcm_harness.paper_revision import RevisionController
    from cumcm_harness.providers import FixtureProvider
    from cumcm_harness.common import tree_manifest,IntegrityError
    out.mkdir(parents=True,exist_ok=True)
    source=out/'source.docx';doc=Document()
    section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Cm(2.5)
    for name in ('Normal','Title','Heading 1'):
        style=doc.styles[name];style.font.name='Noto Serif CJK SC';style.font.size=Pt(11 if name=='Normal' else 16)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    doc.add_paragraph('论文修订对象保留测试',style='Title')
    doc.add_paragraph('本文件用于检查语言修订后，公式、表格、图片、编号、脚注及参考标记的保留情况。')
    doc.add_paragraph('本文模型的误差为 0.25，现有结论仍需按原始证据核验。')
    p=doc.add_paragraph();math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='x = 1';r.append(t);math.append(r);p._p.append(math)
    table=doc.add_table(rows=2,cols=2);table.style='Table Grid'
    table.cell(0,0).text='量';table.cell(0,1).text='原值';table.cell(1,0).text='误差';table.cell(1,1).text='0.25'
    png=io.BytesIO();Image.new('RGB',(80,40),(20,85,120)).save(png,format='PNG');png.seek(0)
    doc.add_picture(png,width=Cm(1.2),height=Cm(.6))
    doc.add_paragraph('保持原有编号和交叉引用[1]。',style='List Number')
    foot=doc.add_paragraph('此段保留原脚注及技术含义。');r=foot.add_run()._r
    marker=OxmlElement('w:footnoteReference');marker.set(qn('w:id'),'2');r.append(marker)
    doc.add_paragraph('参考文献[1]：文档保留测试条目，不代表已核验的科研文献。')
    section.footer.paragraphs[0].text='文档对象保留测试'
    doc.save(source)
    # Add an actual related footnote part, not a textual footnote imitation.
    with zipfile.ZipFile(source) as z:members={i.filename:z.read(i) for i in z.infolist()}
    notes=ET.Element('{'+W+'}footnotes',nsmap={'w':W})
    note=ET.SubElement(notes,'{'+W+'}footnote');note.set('{'+W+'}id','2')
    ET.SubElement(ET.SubElement(ET.SubElement(note,'{'+W+'}p'),'{'+W+'}r'),'{'+W+'}t').text='原脚注内容保持不变。'
    members['word/footnotes.xml']=ET.tostring(notes,xml_declaration=True,encoding='utf-8')
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    rels=ET.fromstring(members['word/_rels/document.xml.rels'])
    ET.SubElement(rels,'{'+relns+'}Relationship',Id='rIdRegressionFootnote',
        Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes',Target='footnotes.xml')
    members['word/_rels/document.xml.rels']=ET.tostring(rels,xml_declaration=True,encoding='utf-8')
    types=ET.fromstring(members['[Content_Types].xml'])
    ET.SubElement(types,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',
        PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    members['[Content_Types].xml']=ET.tostring(types,xml_declaration=True,encoding='utf-8')
    def save_members(path,values):
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in values.items():z.writestr(name,data)
    save_members(source,members)

    def response(role,schema,packet):
        if schema=='review':return {'target_digest':packet['target_digest'],'verdict':'PASS',
            'scope':'Fixed editorial fixture only, not scientific validation.','findings':[],
            'evidence':['Checked by this fixed test responder only.'],'unverified':[]}
        assert schema=='revision_patch'
        return {'source_sha256':packet['source_sha256'],'diagnosis':[{'location':'普通正文','issue':'连接词调整。','scope':'LANGUAGE'}],
            'edits':[{'block_id':b['id'],'before_sha256':b['sha256'],
                'replacement':b['text'].replace('本文模型','本文所用模型'),
                'reason':'只调整普通语句，保留数值、符号和原有结论。'} for b in packet['blocks'] if '本文模型' in b['text']],
            'research_requests':[],'limitations':['未重跑原实验；审查回复为固定工程夹具。']}

    def render(path,label):
        folder=out/('render-'+label);folder.mkdir()
        renderer=input_dir/'render_docx.py'
        if renderer.exists():
            cmd=[sys.executable,str(renderer),str(path),'--output_dir',str(folder),'--emit_pdf']
        else:
            profile=Path('/tmp')/('lo-'+label)
            cmd=['/usr/bin/soffice','-env:UserInstallation='+profile.as_uri(),
                 '--headless','--convert-to','pdf','--outdir',str(folder),str(path)]
        proc=subprocess.run(cmd,capture_output=True,timeout=120)
        (folder/'render.stdout.log').write_bytes(proc.stdout);(folder/'render.stderr.log').write_bytes(proc.stderr)
        if proc.returncode:raise IntegrityError('DOCX renderer failed; inspect its retained logs')
        pdfs=list(folder.glob('*.pdf'))
        if len(pdfs)!=1:raise IntegrityError('Renderer did not produce exactly one PDF')
        document=fitz.open(pdfs[0]);pages=[]
        for i,page in enumerate(document):
            target=folder/f'page-{i+1:03d}.png';page.get_pixmap(matrix=fitz.Matrix(1.4,1.4)).save(target);pages.append(target.name)
        if not pages:raise IntegrityError('Empty rendered document')
        text='\n'.join(p.get_text() for p in document)
        document.close()
        # LibreOffice Writer without its Math component silently drops OMML.
        # The controlled corpus must visibly render the preserved equation.
        if label.startswith('controlled') and not all(token in text for token in ('x', '=', '1')):
            raise IntegrityError('Controlled OMML equation missing from rendered PDF; install LibreOffice Math')
        return {'pages':len(pages),'pdf':str(pdfs[0].relative_to(out)),'text_sha256':sha_text(text),
                'renderer':'PACKAGED_RENDER_DOCX_IN_DOCKER' if renderer.exists() else 'ISOLATED_LIBREOFFICE_AND_PYMUPDF',
                'visual_review':'REQUIRED_SEPARATELY'}

    results=[]
    inputs=[('controlled',source)]+([('supplied',input_dir/'supplied.docx')] if (input_dir/'supplied.docx').exists() else [])
    for label,path in inputs:
        source_hash=file_hash(path);run=out/('revision-'+label)
        initialize(run,input_mode='revise',paper=path,config={**DEFAULT_CONFIG,'materials_workflow':True})
        controller=RevisionController(run,fixture_provider=FixtureProvider(response));result=controller.run()
        revised=run/'deliverables/revised.docx'
        with zipfile.ZipFile(path) as old,zipfile.ZipFile(revised) as new:
            if old.namelist()!=new.namelist():raise IntegrityError('DOCX package inventory changed')
            preserved=[n for n in old.namelist() if n!='word/document.xml']
            if any(old.read(n)!=new.read(n) for n in preserved):raise IntegrityError('Non-body DOCX member changed')
            before=ET.fromstring(old.read('word/document.xml'));after=ET.fromstring(new.read('word/document.xml'))
            for tag in ('{'+M+'}oMath','{'+W+'}tbl','{'+W+'}drawing','{'+W+'}numPr','{'+W+'}footnoteReference'):
                if [ET.tostring(e) for e in before.iter(tag)]!=[ET.tostring(e) for e in after.iter(tag)]:
                    raise IntegrityError('Protected DOCX object changed: '+tag)
        assert source_hash==file_hash(path)
        calls=controller.store.get('model_calls_reserved');manifest=tree_manifest(run/'deliverables')
        def never(*args,**kwargs):raise AssertionError('Editorial replay invoked a new model')
        RevisionController(run,fixture_provider=FixtureProvider(never)).run()
        assert tree_manifest(run/'deliverables')==manifest and controller.store.get('model_calls_reserved')==calls
        results.append({'source':label,'status':'PASS','source_sha256':source_hash,'edits':result['patches'],
            'preserved_other_members':len(preserved),'source_unchanged':True,'replay_unchanged':True,
            'before_render':render(path,label+'-before'),'after_render':render(revised,label+'-after')})

    rejected=[]
    for kind in ('commentRangeStart','ins'):
        values=dict(members);root=ET.fromstring(values['word/document.xml']);p=root.find('.//{'+W+'}p')
        ET.SubElement(p,'{'+W+'}'+kind);values['word/document.xml']=ET.tostring(root)
        path=out/(kind+'.docx');save_members(path,values);original=file_hash(path)
        d=read_document(path,purpose='paper');assert not any(b['editable'] for b in d['blocks'])
        run=out/('reject-'+kind);initialize(run,input_mode='revise',paper=path,config=dict(DEFAULT_CONFIG))
        c=RevisionController(run,fixture_provider=FixtureProvider(never))
        try:c.run()
        except Blocked:pass
        else:raise AssertionError('Review-marked DOCX was edited')
        assert c.store.get('model_calls_reserved',0)==0 and file_hash(path)==original
        rejected.append({'case':kind,'status':'EXPECTED_BLOCK_BEFORE_MODEL','original_unchanged':True})
    summary={'status':'PASS','scope':'FIXTURE_SEMANTICS_WITH_REAL_DOCX_AND_RENDERING_IN_DOCKER',
        'documents':results,'review_markup_negatives':rejected,'original_science_revalidated':False,
        'real_model_calls':0,'visual_review':'REQUIRED_SEPARATELY'}
    write_json(out/'summary.json',summary)


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--out',type=Path,required=True)
    p.add_argument('--source',type=Path);p.add_argument('--renderer',type=Path)
    p.add_argument('--image',default='cumcm-egoharness:0.5.0-rc2-documents')
    p.add_argument('--input',type=Path);p.add_argument('--worker',action='store_true')
    a=p.parse_args()
    if a.worker:worker(a.input,a.out);return
    out=a.out.resolve()
    if out.exists():raise Blocked('Use a fresh document validation directory')
    out.mkdir(parents=True);code=out/'code';code.mkdir();data=out/'inputs';data.mkdir()
    shutil.copy2(__file__,code/'validate.py')
    for source,name in ((a.source,'supplied.docx'),(a.renderer,'render_docx.py')):
        if source:shutil.copy2(source,data/name)
    executor=Executor(image=a.image)
    receipt=executor.execute(code,'validate.py',data,out/'output',['--worker'],
        Limits(seconds=360,memory_mb=2048),logdir=out/'logs')
    summary=read_json(out/'output/summary.json')
    summary['execution_receipt']=receipt;summary['runtime']=executor.probe()
    summary['input_hashes']={p.name:file_hash(p) for p in data.iterdir()}
    write_json(out/'summary.json',summary);print('PASS:',len(summary['documents']),'DOCX pairs rendered')


if __name__=='__main__':main()
